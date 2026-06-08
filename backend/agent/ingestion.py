import logging
import os
from sqlalchemy import text

logger = logging.getLogger(__name__)

# ── Extension-based type buckets ─────────────────────────────────────────────

_PLAIN_TEXT_EXTS = {
    ".txt", ".md", ".rst", ".csv", ".log", ".srt", ".vtt",
    ".ini", ".cfg", ".conf", ".toml",
}

_CODE_EXTS = {
    ".py", ".js", ".ts", ".jsx", ".tsx", ".json", ".yaml", ".yml",
    ".xml", ".html", ".htm", ".css", ".scss", ".sass", ".less",
    ".sh", ".bash", ".zsh", ".fish", ".ps1", ".bat", ".cmd",
    ".c", ".cpp", ".h", ".hpp", ".java", ".go", ".rs", ".rb",
    ".php", ".swift", ".kt", ".r", ".m", ".sql",
}

_MEDIA_EXTS = {
    ".mp4", ".mkv", ".avi", ".mov", ".wmv", ".flv", ".webm",
    ".mp3", ".aac", ".ogg", ".wav", ".flac", ".m4a",
    ".jpg", ".jpeg", ".png", ".gif", ".bmp", ".webp", ".svg", ".tiff",
}


def _extract_text(real_path: str, mime_type: str, filename: str) -> str:
    """
    Extract up to ~500 words from a file for FTS5 indexing.
    Returns "" on any failure or for binary/media types.
    Never raises.
    """
    ext = os.path.splitext(filename)[1].lower()
    mime = (mime_type or "").lower()

    if not real_path or not os.path.exists(real_path):
        return ""

    try:
        # ── PDF ──────────────────────────────────────────────────────────────
        if mime == "application/pdf" or ext == ".pdf":
            try:
                import pdfplumber
                with pdfplumber.open(real_path) as pdf:
                    parts = []
                    word_count = 0
                    for page in pdf.pages:
                        page_text = page.extract_text() or ""
                        parts.append(page_text)
                        word_count += len(page_text.split())
                        if word_count >= 500:
                            break
                    combined = " ".join(parts)
                    return " ".join(combined.split()[:500])
            except Exception as e:
                logger.debug("pdfplumber failed for %s: %s", real_path, e)
                return ""

        # ── Word documents (.docx) ────────────────────────────────────────────
        if (
            ext == ".docx"
            or mime == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ):
            try:
                from docx import Document
                doc = Document(real_path)
                words = []
                for para in doc.paragraphs:
                    words.extend(para.text.split())
                    if len(words) >= 500:
                        break
                return " ".join(words[:500])
            except ImportError:
                logger.debug("python-docx not installed — skipping .docx extraction")
                return ""
            except Exception as e:
                logger.debug("docx extraction failed for %s: %s", real_path, e)
                return ""
        
        # ── PowerPoint (.pptx) ────────────────────────────────────────────────
        if ext == ".pptx" or "presentationml.presentation" in mime:
            try:
                from pptx import Presentation
                prs = Presentation(real_path)
                words = []
                for slide in prs.slides:
                    for shape in slide.shapes:
                        if hasattr(shape, "text"):
                            words.extend(shape.text.split())
                    if len(words) >= 500:
                        break
                return " ".join(words[:500])
            except Exception as e:
                logger.debug("pptx extraction failed for %s: %s", real_path, e)
                return ""
                
        # ── Excel (.xlsx) ────────────────────────────────────────────────────
        if ext == ".xlsx" or "spreadsheetml.sheet" in mime:
            try:
                import openpyxl
                wb = openpyxl.load_workbook(real_path, data_only=True, read_only=True)
                words = []
                for sheet in wb.worksheets:
                    for row in sheet.iter_rows(values_only=True):
                        for cell in row:
                            if cell is not None:
                                words.extend(str(cell).split())
                        if len(words) >= 500:
                            break
                    if len(words) >= 500:
                        break
                return " ".join(words[:500])
            except Exception as e:
                logger.debug("xlsx extraction failed for %s: %s", real_path, e)
                return ""

        # ── Plain text files ──────────────────────────────────────────────────
        if ext in _PLAIN_TEXT_EXTS or mime.startswith("text/"):
            with open(real_path, "r", encoding="utf-8", errors="ignore") as fh:
                content = fh.read(10_000)
            return " ".join(content.split()[:500])

        # ── Source code / structured text files ──────────────────────────────
        if ext in _CODE_EXTS:
            with open(real_path, "r", encoding="utf-8", errors="ignore") as fh:
                content = fh.read(6_000)
            return " ".join(content.split()[:300])

        # ── Media / binary — no text extraction ──────────────────────────────
        # (still indexed with filename + path + mime below)
        return ""

    except Exception as e:
        logger.debug("Text extraction failed for %s: %s", real_path, e)
        return ""


def index_file(
    db_session,          # SQLAlchemy Session or None
    file_id: str,
    real_path: str,
    filename: str,
    relative_path: str,
    mime_type: str,
):
    """
    Insert or replace a row in file_fts for full-text search.

    Works for ALL file types — extracted_text may be "" for media/binary
    files, but filename + relative_path + mime_type are always indexed so
    every file is findable by name even without text content.

    Accepts an existing db_session (SQLAlchemy Session) or None.
    When None, opens its own session via SessionLocal.
    """
    extracted_text = _extract_text(real_path, mime_type, filename)

    sql = text("""
        INSERT OR REPLACE INTO file_fts
            (file_id, filename, relative_path, mime_type, extracted_text)
        VALUES
            (:file_id, :filename, :relative_path, :mime_type, :extracted_text)
    """)
    params = {
        "file_id": file_id,
        "filename": filename,
        "relative_path": relative_path or "",
        "mime_type": mime_type or "",
        "extracted_text": extracted_text,
    }

    if db_session is not None:
        try:
            db_session.execute(sql, params)
            db_session.commit()
            logger.debug("FTS indexed (via session): %s", filename)
        except Exception as e:
            logger.error("FTS index failed for %s: %s", filename, e)
    else:
        try:
            from database import SessionLocal
            with SessionLocal() as db:
                db.execute(sql, params)
                db.commit()
            logger.debug("FTS indexed (own session): %s", filename)
        except Exception as e:
            logger.error("FTS index failed for %s: %s", filename, e)


# ── Compatibility stubs ───────────────────────────────────────────────────────

async def ingest_file(*args, **kwargs):
    """Stub — kept for compatibility if anything still imports it."""
    pass


async def process_queue():
    """Stub — kept for compatibility; main.py may import this."""
    pass
